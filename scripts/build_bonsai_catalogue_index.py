#!/usr/bin/env python3
"""Build the compact, browser-friendly BONSAI v2.1.6 GWP100 catalogue.

The source workbook remains the authority. The generated JSON is deliberately
an index for selecting a background provider in the dataset builder; it does
not turn the static GWP100 lookup into an ontology LCA result.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


BONSAI = "https://w3id.org/mintjesba/pefdpp/bonsai#"


def activity_keys(path: Path) -> set[tuple[str, str]]:
    """Read only the concrete activity identifiers from the published TTL."""
    if not path.exists():
        return set()
    pattern = re.compile(r"^bon:ACT_(.+)__([^\s]+) a act:Activity ;")
    keys = set()
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            match = pattern.match(line)
            if match:
                keys.add((match.group(1), match.group(2)))
    return keys


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--bonsai-co2", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    code_sheet = workbook["flow_codes"]
    footprint_sheet = workbook["footprints-v2.1.6"]

    code_rows = code_sheet.iter_rows(values_only=True)
    next(code_rows)
    codes = {
        str(row[0]): {
            # Nine published rows carry a numeric zero instead of a description.
            # Keep them visible rather than silently dropping valid BONSAI codes.
            "description": str(row[1]) if isinstance(row[1], str) and row[1].strip() else f"BONSAI provider {row[0]} (no published description)",
            "className": str(row[6]),
        }
        for row in code_rows
        if row[0] and row[6]
    }
    concrete_activities = activity_keys(args.bonsai_co2)

    records: dict[str, list[list[object]]] = defaultdict(list)
    footprint_rows = footprint_sheet.iter_rows(values_only=True)
    next(footprint_rows)
    for row in footprint_rows:
        _, code, _, region, reference_unit, value, emission_unit, metric = row
        if metric != "GWP100" or code not in codes:
            continue
        code = str(code)
        region = str(region)
        records[code].append([
            region,
            str(reference_unit),
            float(value),
            str(emission_unit),
            f"{BONSAI}ACT_{code}__{region}" if (code, region) in concrete_activities else "",
        ])

    entries = []
    total_records = 0
    mapped_records = 0
    for code in sorted(codes, key=lambda value: (codes[value]["description"].lower(), value.lower())):
        offerings = sorted(records.get(code, []), key=lambda item: (item[0], item[1]))
        total_records += len(offerings)
        mapped_records += sum(bool(item[4]) for item in offerings)
        entries.append([code, codes[code]["description"], codes[code]["className"], offerings])

    payload = {
        "metadata": {
            "source": "bonsai.xlsx / footprints-v2.1.6",
            "version": "v2.1.6",
            "metric": "GWP100",
            "description": "Static BONSAI GWP100 lookup. Values are catalogue intensities, not calculated PEFStudy results.",
            "bonsaiNamespace": BONSAI,
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "providerCodes": len(entries),
            "gwp100Records": total_records,
            "recordsWithPublishedActivityUri": mapped_records,
        },
        "entries": entries,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(entries)} providers and {total_records} GWP100 records to {args.output}")


if __name__ == "__main__":
    main()
