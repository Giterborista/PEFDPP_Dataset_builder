#!/usr/bin/env python3
"""Build an exact-EF-UUID climate factor index from the supplied EF 3.1 workbook.

No name matching, compartment substitution, direction inversion, unit guessing,
or summation of climate subcategories is performed.
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import openpyxl

METHOD_UUID = '6209b35f-9447-40b5-b68c-a1099e3674a0'

def build(workbook_path, elementary_index_path, destination):
    source = Path(workbook_path)
    elementary = json.loads(Path(elementary_index_path).read_text())
    columns = elementary['metadata']['columns']
    entries = [dict(zip(columns, row)) for row in elementary['entries']]
    units_by_id = {e['identifier']: e['preferredUnitLocal'] for e in entries}
    uuid_by_local = {e['uriLocal']: e['identifier'] for e in entries}
    assert len(units_by_id) == len(entries), 'Duplicate EF flow UUIDs'
    classifications = {e['identifier']: e['classificationLabel'].lower() for e in entries}
    known_by_unit = defaultdict(list)
    for local, identifier in uuid_by_local.items():
        assert local == identifier, 'Index currently expects exact UUID local names'
        classification = classifications[identifier]
        direction = 'output' if 'emissions' in classification else 'input' if 'resources' in classification else 'either'
        known_by_unit[units_by_id[identifier]+'|'+direction].append(identifier)
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    methods = [r for r in book['lciamethods'].iter_rows(min_row=2, values_only=True) if r[0] == METHOD_UUID]
    assert len(methods) == 1 and methods[0][1] == 'Climate change'
    method = methods[0]
    factors = defaultdict(list)
    seen = set()
    directions = Counter()
    locations = Counter()
    for row in book['lciamethods_CF'].iter_rows(min_row=2, values_only=True):
        if row[2] != METHOD_UUID:
            continue
        flow_id, label, method_id, method_name, factor, location, _, _, compartment, derivation, direction = row
        assert method_name == 'Climate change', row
        assert isinstance(factor, (int, float)), row
        assert direction in ('Input', 'Output'), row
        assert flow_id in units_by_id, ('CF outside published elementary-flow index', row)
        region = str(location or 'GLO')
        key = (flow_id, region, direction)
        assert key not in seen, ('Duplicate factor for exact UUID/region/direction', row)
        seen.add(key)
        factors[flow_id].append([region, direction.lower(), factor])
        locations[region] += 1
        directions[direction] += 1
    book.close()
    result = {
        'metadata': {
            'generatedAt': datetime.now(timezone.utc).isoformat(),
            'method': {'name': method[1], 'uuid': METHOD_UUID, 'indicator': method[3], 'year': method[7], 'version': method[8]},
            'source': {'filename': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheet': 'lciamethods_CF', 'officialInformation': 'https://eplca.jrc.ec.europa.eu/LCDN/developerEF.xhtml'},
            'elementaryIndexSha256': hashlib.sha256(Path(elementary_index_path).read_bytes()).hexdigest(),
            'knownFlowCount': len(entries), 'factorCount': len(seen), 'characterisedFlowCount': len(factors),
            'columns': ['region', 'direction', 'kgCO2ePerReferenceUnit'],
            'rule': 'Exact EF 3.1 flow UUID, native reference unit and exchange direction. Use the activity geography when an exact regional factor exists, otherwise the published GLO factor only. Known EF 3.1 flows absent from this complete climate-method CF table have no factor under this method; their zero method contribution is disclosed, not interpreted as zero overall environmental impact. Unknown UUIDs, unsupported directions, incompatible units and missing regional factors block a complete result. Climate subcategories are not added to the total method.',
            'locations': dict(locations), 'directions': dict(directions)
        },
        'knownFlowsByUnitAndDirection': {u: sorted(ids) for u, ids in sorted(known_by_unit.items())},
        'factors': dict(sorted(factors.items()))
    }
    Path(destination).write_text(json.dumps(result, separators=(',', ':'), ensure_ascii=False))
    print(json.dumps(result['metadata'], indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--elementary-index', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.workbook, args.elementary_index, args.output)
