#!/usr/bin/env python3
"""Convert the PV LCI workbook into PEFDPP-compatible Turtle and JSON-LD."""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE = "https://w3id.org/cerise-see/examples/pv-lci#"
DATASET = "https://w3id.org/mintjesba/pefdpp/dataset#"
ACT = "https://w3id.org/mintjesba/pefdpp/activity#"
FLOW = "https://w3id.org/mintjesba/pefdpp/flow#"
PRS = "https://w3id.org/mintjesba/pefdpp/product-system#"
OM = "http://www.ontology-of-units-of-measure.org/resource/om-2/"
DCTERMS = "http://purl.org/dc/terms/"
RDFS = "http://www.w3.org/2000/01/rdf-schema#"
XSD = "http://www.w3.org/2001/XMLSchema#"
EF_BASE = "https://w3id.org/mintjesba/pefdpp/skos/ef31-elementaryflows#"
GEO_BASE = "https://w3id.org/mintjesba/pefdpp/skos/pef-geographies#"
UNIT_BASE = "https://w3id.org/mintjesba/pefdpp/ef31-units#"

UNIT_MAP = {
    "kg": (f"{UNIT_BASE}Kilogram", 1.0, "kg"),
    "m2": (f"{UNIT_BASE}SquareMetre", 1.0, "m2"),
    "m3": (f"{UNIT_BASE}CubicMetre", 1.0, "m3"),
    "MJ": (f"{UNIT_BASE}Megajoule", 1.0, "MJ"),
    "kWh": (f"{UNIT_BASE}Megajoule", 3.6, "MJ"),
    "tkm": (f"{UNIT_BASE}TonneKilometre", 1.0, "tkm"),
    "unit": (f"{UNIT_BASE}Item", 1.0, "item(s)"),
}

GEO_ALIASES = {"APAC": "RAS"}


def life_cycle_stage_for_sheet(sheet_name: str) -> str:
    raw_material_prefixes = ("X-MG-Si", "X-SoG", "X-Si-Market")
    local = "RawMaterialAcquisition" if sheet_name.startswith(raw_material_prefixes) else "Manufacturing"
    return f"{PRS}{local}"

COMPARTMENTS = {
    "air, unspecified": "Emissions to air, unspecified",
    "air, low population density": "Emissions to non-urban air or from high stacks",
    "air, high population density": "Emissions to urban air close to ground",
    "water, river": "Emissions to fresh water",
    "resource, in water": "Renewable material resources from water",
}

FLOW_ALIASES = {
    "heat, waste": "waste heat",
    "carbon monoxide, biogenic": "carbon monoxide (biogenic)",
    "carbon monoxide, fossil": "carbon monoxide (fossil)",
    "carbon dioxide, biogenic": "carbon dioxide (biogenic)",
    "carbon dioxide, fossil": "carbon dioxide (fossil)",
    "nmvoc, non-methane volatile organic compounds, unspecified origin": "non-methane volatile organic compounds",
    "aox, adsorbable organic halogen as cl": "adsorbable organic halogen compounds",
    "bod5, biological oxygen demand": "biological oxygen demand",
    "cod, chemical oxygen demand": "chemical oxygen demand",
    "toc, total organic carbon": "total organic carbon",
    "water, cn": "water vapour",
    "water, cooling, unspecified natural origin, cn": "freshwater",
    "2-propanol": "isopropanol",
    "ethane, 1,1,1,2-tetrafluoro-, hfc-134a": "hfc-134a",
    "particulates, > 10 um": "particles (> pm10)",
    "nitrogen": "nitrogen, total (excluding n2)",
}


def slug(value: str) -> str:
    text = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return text or "item"


def ttl_text(value: object) -> str:
    return str(value).replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")


def number_text(value: object) -> str:
    return format(float(value), ".15g")


def load_indexes(options_path: Path, elementary_path: Path):
    options = json.loads(options_path.read_text(encoding="utf-8"))
    geographies = {item["identifier"]: item for item in options["geographies"]}
    elementary = json.loads(elementary_path.read_text(encoding="utf-8"))
    columns = elementary["metadata"]["columns"]
    by_label: dict[str, list[dict]] = defaultdict(list)
    for packed in elementary["entries"]:
        row = dict(zip(columns, packed))
        by_label[row["label"].strip().lower()].append(row)
    return geographies, by_label


def elementary_match(description: str, compartment: str, by_label: dict[str, list[dict]]):
    label = FLOW_ALIASES.get(description.strip().lower(), description.strip().lower())
    classification = COMPARTMENTS.get(compartment.strip().lower())
    matches = [
        row for row in by_label.get(label, [])
        if row["classificationLabel"].strip().lower() == (classification or "").lower()
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Elementary-flow mapping is not unique for {description!r} in {compartment!r}: {len(matches)} matches"
        )
    return matches[0]


def convert(input_path: Path, output_dir: Path, options_path: Path, elementary_path: Path):
    geographies, elementary_by_label = load_indexes(options_path, elementary_path)
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    dataset_uri = f"{BASE}pv-supply-chain-lci"
    activities = []
    flow_types: dict[str, dict] = {}
    conversion_rows = []
    elementary_count = 0
    kwh_count = 0
    sd95_omitted = 0

    def product_flow_type(description: str, location: str):
        key = f"{description.strip()}|{location.strip()}"
        uri = f"{BASE}flowtype-{slug(description)}-{slug(location)}"
        flow_types.setdefault(uri, {
            "uri": uri,
            "label": f"{description.strip()} — {location.strip()}",
            "class": "ProductFlow",
            "sourceKey": key,
        })
        return uri

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        rows = [row for row in rows if row[0] and row[1] and row[4] is not None]
        product_rows = [row for row in rows if str(row[0]).strip().lower() == "product"]
        if len(product_rows) != 1:
            raise ValueError(f"{sheet.title}: expected exactly one product row, found {len(product_rows)}")
        product = product_rows[0]
        activity_geo = GEO_ALIASES.get(str(product[2]).strip(), str(product[2]).strip())
        if activity_geo not in geographies:
            raise ValueError(f"{sheet.title}: geography {activity_geo!r} is not in the ontology geography index")
        activity_id = f"activity-{slug(sheet.title)}"
        activity_uri = f"{BASE}{activity_id}"
        flows = []
        for row_number, row in enumerate(rows, start=2):
            row_type, description, location, source_unit, source_amount, sd95 = row[:6]
            row_type = str(row_type).strip().lower()
            description = str(description).strip()
            location = str(location).strip()
            source_unit = str(source_unit).strip()
            if source_unit not in UNIT_MAP:
                raise ValueError(f"{sheet.title}!D{row_number}: unsupported unit {source_unit!r}")
            unit_uri, factor, normalized_unit = UNIT_MAP[source_unit]
            amount = float(source_amount) * factor
            if source_unit == "kWh":
                kwh_count += 1
            if sd95 is not None:
                sd95_omitted += 1
            direction = "output" if row_type in {"product", "direct emission"} else "input"
            determining = row_type == "product"
            if row_type == "direct emission":
                match = elementary_match(description, location, elementary_by_label)
                flow_type_uri = f"{EF_BASE}{match['uriLocal']}"
                direction = "input" if location.lower().startswith("resource") else "output"
                if unit_uri != f"{UNIT_BASE}{match['preferredUnitLocal']}":
                    raise ValueError(
                        f"{sheet.title}!D{row_number}: unit {unit_uri} differs from elementary-flow preferred unit {match['preferredUnitLocal']}"
                    )
                ontology_label = match["label"]
                elementary_count += 1
            elif row_type in {"product", "component", "resource"}:
                flow_type_uri = product_flow_type(description, location)
                ontology_label = description
            else:
                raise ValueError(f"{sheet.title}!A{row_number}: unsupported row type {row_type!r}")
            flow_uri = f"{BASE}flow-{slug(sheet.title)}-{row_number}"
            flows.append({
                "uri": flow_uri,
                "label": f"{description} [{location}]",
                "direction": direction,
                "flowTypeUri": flow_type_uri,
                "amount": number_text(amount),
                "unitUri": unit_uri,
                "determining": determining,
            })
            conversion_rows.append({
                "sheet": sheet.title,
                "excelRow": row_number,
                "sourceType": row_type,
                "sourceDescription": description,
                "sourceLocation": location,
                "sourceUnit": source_unit,
                "sourceAmount": source_amount,
                "ontologyFlowClass": "ElementaryFlow" if row_type == "direct emission" else "ProductFlow",
                "ontologyFlowLabel": ontology_label,
                "flowTypeUri": flow_type_uri,
                "direction": direction,
                "normalizedUnit": normalized_unit,
                "unitUri": unit_uri,
                "normalizedAmount": number_text(amount),
                "conversion": "kWh × 3.6 = MJ" if source_unit == "kWh" else "none",
            })
        activities.append({
            "uri": activity_uri,
            "name": str(product[1]).strip(),
            "activityLinkUri": f"{BASE}activity-link-{slug(sheet.title)}",
            "lifeCycleStageUri": life_cycle_stage_for_sheet(sheet.title),
            "geographyUri": geographies[activity_geo]["uri"],
            "companySpecific": True,
            "flows": flows,
        })

    graph = [{
        "@id": dataset_uri,
        "@type": f"{DATASET}LCIDataset",
        f"{DCTERMS}title": "Photovoltaic panel supply-chain LCI",
        f"{DATASET}hasDatasetType": {"@id": f"{DATASET}DPP"},
        f"{DATASET}hasComplianceLevel": {"@id": f"{DATASET}EFCompliant"},
        f"{DATASET}containsActivity": [{"@id": activity["uri"]} for activity in activities],
    }]
    product_system_uri = f"{BASE}product-system"
    graph.append({
        "@id": product_system_uri,
        "@type": f"{PRS}ProductSystem",
        f"{RDFS}label": "Photovoltaic panel product system",
        f"{PRS}hasActivityLink": [{"@id": activity["activityLinkUri"]} for activity in activities],
    })
    for item in flow_types.values():
        graph.append({
            "@id": item["uri"], "@type": f"{FLOW}{item['class']}",
            f"{FLOW}sourcedFrom": {"@id": dataset_uri}, f"{RDFS}label": item["label"],
        })
    for activity in activities:
        determining = next(flow for flow in activity["flows"] if flow["determining"])
        graph.append({
            "@id": activity["activityLinkUri"], "@type": f"{PRS}ActivityLink",
            f"{PRS}forActivity": {"@id": activity["uri"]},
            f"{PRS}hasLifeCycleStage": {"@id": activity["lifeCycleStageUri"]},
        })
        graph.append({
            "@id": activity["uri"], "@type": f"{ACT}Activity",
            f"{RDFS}label": activity["name"],
            f"{ACT}isCompanySpecific": activity["companySpecific"],
            f"{ACT}hasGeography": {"@id": activity["geographyUri"]},
            f"{ACT}hasDeterminingFlow": {"@id": determining["uri"]},
        })
        for flow in activity["flows"]:
            graph.append({
                "@id": flow["uri"], "@type": f"{FLOW}Flow",
                f"{FLOW}hasFlowType": {"@id": flow["flowTypeUri"]},
                f"{FLOW}contextualLabel": flow["label"],
                f"{ACT}{'outputOf' if flow['direction'] == 'output' else 'inputOf'}": {"@id": activity["uri"]},
                f"{FLOW}hasMeasure": {
                    "@type": f"{OM}Measure",
                    f"{OM}hasUnit": {"@id": flow["unitUri"]},
                    f"{OM}hasNumericalValue": {"@value": flow["amount"], "@type": f"{XSD}decimal"},
                },
            })

    output_dir.mkdir(parents=True, exist_ok=True)
    jsonld_path = output_dir / "pv_panel_lci_ontology.jsonld"
    ttl_path = output_dir / "pv_panel_lci_ontology.ttl"
    report_path = output_dir / "pv_panel_lci_conversion_report.json"
    jsonld_path.write_text(json.dumps({"@context": {"@vocab": FLOW, "dcterms": DCTERMS}, "@graph": graph}, indent=2), encoding="utf-8")

    lines = [
        f"@prefix : <{BASE}> .", f"@prefix dataset: <{DATASET}> .", f"@prefix act: <{ACT}> .",
        f"@prefix flow: <{FLOW}> .", f"@prefix prs: <{PRS}> .", f"@prefix om: <{OM}> .", f"@prefix dcterms: <{DCTERMS}> .",
        f"@prefix rdfs: <{RDFS}> .", f"@prefix xsd: <{XSD}> .", "",
        f"<{dataset_uri}> a dataset:LCIDataset ;",
        '  dcterms:title "Photovoltaic panel supply-chain LCI"@en ;',
        "  dataset:hasDatasetType dataset:DPP ;",
        "  dataset:hasComplianceLevel dataset:EFCompliant ;",
        "  dataset:containsActivity " + ", ".join(f"<{activity['uri']}>" for activity in activities) + " .", "",
        f"<{product_system_uri}> a prs:ProductSystem ;",
        '  rdfs:label "Photovoltaic panel product system"@en ;',
        "  prs:hasActivityLink " + ", ".join(f"<{activity['activityLinkUri']}>" for activity in activities) + " .", "",
    ]
    for item in flow_types.values():
        lines.extend([
            f"<{item['uri']}> a flow:{item['class']} ;",
            f"  flow:sourcedFrom <{dataset_uri}> ;",
            f'  rdfs:label "{ttl_text(item["label"])}"@en .', "",
        ])
    for activity in activities:
        determining = next(flow for flow in activity["flows"] if flow["determining"])
        lines.extend([
            f"<{activity['activityLinkUri']}> a prs:ActivityLink ;",
            f"  prs:forActivity <{activity['uri']}> ;",
            f"  prs:hasLifeCycleStage <{activity['lifeCycleStageUri']}> .", "",
            f"<{activity['uri']}> a act:Activity ;",
            f'  rdfs:label "{ttl_text(activity["name"])}"@en ;',
            f'  act:isCompanySpecific "true"^^xsd:boolean ;',
            f"  act:hasGeography <{activity['geographyUri']}> ;",
            f"  act:hasDeterminingFlow <{determining['uri']}> .", "",
        ])
        for flow in activity["flows"]:
            relation = "act:outputOf" if flow["direction"] == "output" else "act:inputOf"
            lines.extend([
                f"<{flow['uri']}> a flow:Flow ;",
                f"  flow:hasFlowType <{flow['flowTypeUri']}> ;",
                f'  flow:contextualLabel "{ttl_text(flow["label"])}"@en ;',
                f"  {relation} <{activity['uri']}> ;",
                "  flow:hasMeasure [ a om:Measure ;",
                f"    om:hasUnit <{flow['unitUri']}> ;",
                f'    om:hasNumericalValue "{flow["amount"]}"^^xsd:decimal ] .', "",
            ])
    ttl_path.write_text("\n".join(lines), encoding="utf-8")

    report = {
        "sourceWorkbook": str(input_path),
        "outputs": {"turtle": str(ttl_path), "jsonLd": str(jsonld_path)},
        "summary": {
            "activities": len(activities),
            "flows": sum(len(activity["flows"]) for activity in activities),
            "productFlowTypes": len(flow_types),
            "elementaryFlowRowsMappedToExactEf31Uris": elementary_count,
            "kWhRowsConvertedToMJ": kwh_count,
            "sd95ValuesNotExportedBecauseTheCurrentMvpHasNoUncertaintyField": sd95_omitted,
            "bonsaiFieldsWritten": 0,
            "productSystems": 1,
            "activityLinks": len(activities),
        },
        "normalizations": {
            "APAC": "RAS (Asia ontology geography)",
            "kWh": "MJ using amount × 3.6",
            "unit": "item(s)",
            "directElementaryFlows": "Exact EF 3.1 flow URI selected by normalized label and ontology compartment",
            "lifeCycleStages": "MG-Si, SoG-Si and Si market sheets use RawMaterialAcquisition; wafer, cell and panel transformation sheets use Manufacturing",
        },
        "rows": conversion_rows,
    }
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--options", type=Path, required=True)
    parser.add_argument("--elementary-index", type=Path, required=True)
    args = parser.parse_args()
    report = convert(args.input, args.output_dir, args.options, args.elementary_index)
    print(json.dumps(report["summary"], indent=2))


if __name__ == "__main__":
    main()
